"""Register-derived priors: the PERT mapping, the unit reconciliation, and coverage."""

import pytest

from forecast.params import Molecule
from forecast.priors import (

    LAMBDA_DEFAULT,
    LAMBDA_ELICITED,
    MissingRegisterRows,
    load_priors,
    pert,
    pert_mean,
)


@pytest.fixture(scope="module")
def priors():
    # The register is known not to cover every parameter -- see test_coverage_gap.
    return load_priors(require_coverage=False)


# ---- the distribution ---------------------------------------------------------------


@pytest.mark.parametrize(
    "low,base,high,lam",
    [(110, 145, 195, 4.0), (110, 145, 195, 2.0), (0.02, 0.04, 0.07, 4.0),
     (62, 78, 90, 2.0), (3.5, 5, 6.5, 4.0), (2028, 2029, 2031, 4.0)],
)
def test_pert_mean_matches_the_register_formula(low, base, high, lam):
    """mu = (Low + lambda*Base + High) / (lambda + 2)."""
    expected = (low + lam * base + high) / (lam + 2)
    assert pert_mean(low, base, high, lam) == pytest.approx(expected)
    assert pert(low, base, high, lam).mean() == pytest.approx(expected, rel=1e-12)


def test_pert_is_supported_on_low_to_high():
    d = pert(110, 145, 195)
    assert d.ppf(0.0) == pytest.approx(110)
    assert d.ppf(1.0) == pytest.approx(195)


def test_elicited_lambda_is_wider_between_the_same_endpoints():
    """Lower lambda puts less weight on a mode that is expert judgement."""
    standard = pert(62, 78, 90, LAMBDA_DEFAULT)
    elicited = pert(62, 78, 90, LAMBDA_ELICITED)
    assert elicited.std() > standard.std()


def test_degenerate_row_is_a_point_mass():
    assert list(pert(5, 5, 5).rvs(3)) == [5, 5, 5]


@pytest.mark.parametrize("args", [(1, 0, 2), (0, 3, 2), (2, 1, 0)])
def test_pert_rejects_a_base_outside_the_range(args):
    with pytest.raises(ValueError, match="low <= base <= high"):
        pert(*args)


def test_pert_rejects_a_non_positive_shape():
    with pytest.raises(ValueError, match="shape must be positive"):
        pert(1, 2, 3, lam=0)


# ---- reading the register -----------------------------------------------------------


def test_lambda_is_set_by_source_type(priors):
    for group in priors.by_param.values():
        for p in group:
            want = LAMBDA_ELICITED if p.source_type == "Elicited" else LAMBDA_DEFAULT
            assert p.lam == want, f"{p.param} (row {p.row_id}, {p.source_type})"


def test_elicited_rows_are_actually_present(priors):
    """Guards the test above from passing vacuously."""
    elicited = [p for g in priors.by_param.values() for p in g if p.lam == LAMBDA_ELICITED]
    assert {p.param for p in elicited} == {"remission", "brandAttr", "wac"}


def test_derived_rows_are_skipped(priors):
    assert "peakShare" not in priors.by_param
    skipped = {s.param: s for s in priors.skipped}
    assert "peakShare" in skipped
    assert skipped["peakShare"].source_type == "Derived"


def test_rows_with_blank_estimates_are_skipped_not_guessed(priors):
    """competitors.attr is registered but never filled in."""
    assert "competitors.attr" not in priors.by_param
    reasons = {s.param: s.reason for s in priors.skipped}
    assert "blank" in reasons["competitors.attr"]


def test_every_row_is_either_a_prior_or_an_explained_skip(priors):
    """Nothing is silently dropped: every data row becomes a prior or a stated skip.

    Counted against the workbook itself rather than a hard-coded total, so editing the
    register cannot quietly invalidate the check.
    """
    import openpyxl

    from forecast.priors import HEADER_ROW, SHEET

    ws = openpyxl.load_workbook(priors.register_path, data_only=True)[SHEET]
    data_rows = sum(
        1
        for r in range(HEADER_ROW + 1, ws.max_row + 1)
        if isinstance(ws.cell(r, 1).value, (int, float))
    )
    assert data_rows > 0
    assert len(priors) + len(priors.skipped) == data_rows
    assert all(s.reason for s in priors.skipped)


# ---- the unit reconciliation --------------------------------------------------------


def test_percent_rows_are_converted_to_the_model_convention(priors):
    """The register stores 0.28; the model stores 28 (CLAUDE.md)."""
    for group in priors.by_param.values():
        for p in group:
            want = 100.0 if "%" in p.unit else 1.0
            assert p.scaled_by == want, f"{p.param} has unit {p.unit!r}"


@pytest.mark.parametrize(
    "param,expected",
    [
        # Percent rows, stored as fractions, that land exactly on the example
        # parameters once multiplied by 100. This is what pins the conversion rule.
        ("orderPenalty", 12.0),
        ("freeGoodsPct", 3.0),
        ("returnsPct", 1.2),
        ("feeWholesale", 4.5),
        ("feeSP", 2.5),
        ("discountRate", 9.0),
        # Non-percent rows, which must pass through untouched.
        ("populationM", 340.0),
        ("bassP", 0.04),
    ],
)
def test_converted_base_matches_the_example_parameters(priors, param, expected):
    assert priors.one(param).base == pytest.approx(expected)


def test_converted_share_percentages_are_plausible(priors):
    """A missed conversion shows up as a share-like percentage sitting below 1.

    Annual rates are excluded: population growth is legitimately 0.3-0.7 %/yr and price
    growth can be negative, so neither can be judged by magnitude.
    """
    for group in priors.by_param.values():
        for p in group:
            if "%" in p.unit and p.unit != "% per year":
                assert p.high > 1.0, f"{p.param} looks unconverted: high={p.high}"


def test_annual_rates_are_converted_too(priors):
    """The rows excluded above are still scaled -- just not checkable by magnitude."""
    assert priors.one("popGrowth").base == pytest.approx(0.5)  # params default
    assert priors.one("priceGrowth").base == pytest.approx(1.5)  # register base 0.015


# ---- repeated parameters ------------------------------------------------------------


def test_repeated_parameters_keep_every_row(priors):
    """One row per payer segment -- collapsing to one key would discard five of six."""
    assert len(priors.by_param["segments.discount"]) == 6
    assert len(priors.by_param["segments.mix"]) == 6
    assert len(priors.by_param["segments.access"]) == 6
    assert len(priors.by_param["competitors.entry"]) == 2


def test_one_refuses_to_pick_arbitrarily_among_repeats(priors):
    with pytest.raises(KeyError, match="register rows"):
        priors.one("segments.discount")


def test_one_raises_for_an_unregistered_parameter(priors):
    with pytest.raises(KeyError, match="no register row"):
        priors.one("promoIntensity")


# ---- coverage -----------------------------------------------------------------------


# The parameters the register does not cover today. This is a live finding, not a
# permanent fact: as rows are added to the register, remove them from this set. The
# test exists so that closing a gap is a deliberate edit rather than a silent change.
KNOWN_GAPS = {
    # The entire promotional and halo block, which MODEL_DECISIONS sections 4, 6 and 7
    # identify as high leverage -- turning the halo off costs 11.9% of cumulative net.
    "Molecule.promo_intensity",
    "Molecule.sov_theta",
    "Molecule.halo_t",
    "Molecule.halo_a",
    "Molecule.halo_seed",
    # Multi-indication couplings: the register predates the multi-indication model.
    "Indication.prescriber_overlap",
    "Indication.patient_overlap",
    "Indication.sov_share",
    "Indication.launch_year",
    # In-line and loss-of-exclusivity branches, unexercised by the example parameters.
    "Indication.current_share",
    "Indication.share_delta",
    "Indication.loe_year",
    "Indication.loe_erosion",
    "Indication.persist_months",
    # Registered but never filled in.
    "Competitor.attr",
}


def test_coverage_gap_is_exactly_the_known_set():
    priors = load_priors(require_coverage=False)
    assert set(priors.missing) == KNOWN_GAPS


def test_load_priors_fails_loudly_by_default():
    with pytest.raises(MissingRegisterRows) as exc:
        load_priors()
    assert set(exc.value.missing) == KNOWN_GAPS


def test_the_failure_names_the_parameters_and_how_to_proceed():
    with pytest.raises(MissingRegisterRows) as exc:
        load_priors()
    message = str(exc.value)
    for gap in KNOWN_GAPS:
        assert gap in message
    assert "require_coverage=False" in message


def test_an_unfilled_row_is_reported_differently_from_an_absent_one():
    """Competitor.attr has a row; Molecule.halo_t does not. Different fixes."""
    priors = load_priors(require_coverage=False)
    assert "row 19" in priors.missing["Competitor.attr"]
    assert priors.missing["Molecule.halo_t"] == "no register row"


def test_structural_settings_are_not_reported_as_gaps():
    """Mode switches and identifiers are scenario definition, not uncertainty."""
    priors = load_priors(require_coverage=False)
    for exempt in ("Indication.mode", "Indication.epi_mode", "Molecule.horizon",
                   "Indication.name", "Molecule.base_year"):
        assert exempt not in priors.missing


def test_scope_mismatch_is_reported(priors):
    """The register scopes abrasion to the payer segment; the model does not."""
    assert "segments.abrasion" in priors.scope_mismatches


def test_every_prior_points_at_a_real_field(priors):
    """A prior whose field has been renamed or removed is a broken link."""
    models = {
        "Molecule": Molecule.model_fields,
        "Indication": __import__(
            "forecast.params", fromlist=["Indication"]
        ).Indication.model_fields,
        "Segment": __import__(
            "forecast.params", fromlist=["Segment"]
        ).Segment.model_fields,
        "Competitor": __import__(
            "forecast.params", fromlist=["Competitor"]
        ).Competitor.model_fields,
    }
    for group in priors.by_param.values():
        for p in group:
            model_name, _, fname = p.field.partition(".")
            assert fname in models[model_name], f"{p.param} -> {p.field}"
