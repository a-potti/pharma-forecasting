"""Register-derived prior distributions.

Reads an assumptions register and turns each assumption's Low/Base/High into a modified
PERT distribution, keyed by the register's "Model parameter" column. The register is the
source of truth for parameter values (CLAUDE.md), so this module is the only place a
number crosses from the register into the model.

The default is params/example_register.xlsx, a synthetic register committed so the
repository is self-contained -- see its Read me tab. Point ``load_priors(path=...)`` at a
real register to use one; real registers are gitignored because vendor licences prohibit
redistribution and git history is permanent.

Unlike engine.py this module does I/O -- it is the boundary layer that reads the register
so that engine.py never has to.

Modified PERT
-------------
For Low a, Base m, High b and shape lambda::

    mu = (a + lambda*m + b) / (lambda + 2)

which is the mean of a Beta on [a, b] with::

    alpha = 1 + lambda*(m - a)/(b - a)
    beta  = 1 + lambda*(b - m)/(b - a)

lambda = 4 recovers the standard PERT. Rows whose Source type is "Elicited" use
lambda = 2 instead: expert judgement puts less weight on the mode, so the distribution
is deliberately wider between the same endpoints (see the Read me tab, which asks for a
wider range and rarely High confidence on elicited rows).

Units
-----
**The register and the engine disagree, and this module is where that is reconciled.**
CLAUDE.md fixes the model convention as percentages stored as ``28``, not ``0.28``, but
the register stores every percent-unit row as a fraction: ``discountRate`` 0.09,
``copayPct`` 0.045, ``segments.discount`` 0.28. Any row whose Unit column contains "%"
is therefore multiplied by 100 on the way in.

That rule is not a guess. Six rows survive the conversion as exact matches against
params/example_ibd.yaml -- orderPenalty 0.12 -> 12, freeGoodsPct 0.03 -> 3.0,
returnsPct 0.012 -> 1.2, feeWholesale 0.045 -> 4.5, feeSP 0.025 -> 2.5,
discountRate 0.09 -> 9 -- which would be a remarkable coincidence if the rule were wrong.
The same six hold against the real register that this convention was first found in.
``Prior.scaled_by`` records the factor applied to each row so the conversion is auditable
rather than invisible.

Coverage
--------
``load_priors`` raises ``MissingRegisterRows`` when a model parameter has no register row
behind it. Structural choices -- names, mode switches, the horizon -- are exempt, because
they are scenario definition rather than uncertain quantity. Everything else must be
registered before it can be sampled: an unregistered parameter silently held at its
default is a driver the Monte Carlo will report zero uncertainty on.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import openpyxl
from scipy import stats

from .params import Competitor, Indication, Molecule, Segment

__all__ = [
    "Prior",
    "Priors",
    "SkippedRow",
    "MissingRegisterRows",
    "load_priors",
    "pert",
]

# The synthetic example register, committed so the repository is self-contained. A real
# register goes at reference/assumptions_register.xlsx and is loaded with
# load_priors(path=...); it is gitignored because vendor licences prohibit
# redistribution and git history is permanent.
DEFAULT_REGISTER = (
    Path(__file__).resolve().parents[2] / "params" / "example_register.xlsx"
)
REAL_REGISTER = (
    Path(__file__).resolve().parents[2] / "reference" / "assumptions_register.xlsx"
)

SHEET = "Register"
HEADER_ROW = 5

LAMBDA_DEFAULT = 4.0
LAMBDA_ELICITED = 2.0

# Source types that carry no distribution of their own.
SKIP_SOURCE_TYPES = {"Derived"}

# Register "Model parameter" spellings that camel-to-snake cannot reach, or that are
# scoped to the wrong model. Maps the register's name to "Model.field".
EXPLICIT_FIELD_MAP = {
    "feeSP": "Molecule.fee_sp",
    "persist12": "Indication.persist12",
    "competitors.attr": "Competitor.attr",
    "competitors.entry": "Competitor.entry",
    "segments.mix": "Segment.mix",
    "segments.access": "Segment.access",
    "segments.discount": "Segment.discount",
    # The register scopes abrasion to the payer segment. In the multi-indication model
    # it is indication-level: payers list the molecule, but utilisation management is
    # written per indication (docs/MODEL_DECISIONS.md section 5). Mapped to where the
    # parameter actually lives, and reported as a scope mismatch.
    "segments.abrasion": "Indication.abrasion",
}

# Register rows whose scope disagrees with the model, reported alongside the priors.
SCOPE_MISMATCHES = {
    "segments.abrasion": (
        "registered against the payer segment, but abrasion is indication-level in the "
        "multi-indication model (MODEL_DECISIONS section 5)"
    ),
}

# Parameters that are scenario structure rather than uncertain quantities, so they need
# no register row. Each carries the reason it is exempt.
EXEMPT_FIELDS = {
    "Molecule.asset": "asset identity, not a quantity",
    "Molecule.base_year": "scenario definition",
    "Molecule.horizon": "scenario definition",
    "Molecule.segments": "container",
    "Molecule.indications": "container",
    "Segment.name": "identity",
    "Competitor.name": "identity",
    "Indication.name": "identity",
    "Indication.mode": "structural switch (pipeline vs in-line)",
    "Indication.epi_mode": "structural switch (prevalence vs incidence flow)",
    "Indication.persist_mode": "structural switch (curve vs direct)",
    "Indication.comp_on": "structural switch (derive share from competition)",
    "Indication.loe_on": "structural switch (loss of exclusivity modelled)",
    "Indication.competitors": "container",
}

_MODELS = {
    "Molecule": Molecule,
    "Indication": Indication,
    "Segment": Segment,
    "Competitor": Competitor,
}


# ---------------------------------------------------------------------------------
# distribution
# ---------------------------------------------------------------------------------


def pert(low: float, base: float, high: float, lam: float = LAMBDA_DEFAULT):
    """A modified PERT as a frozen scipy distribution on [low, high].

    ``lam`` is the weight on the mode. The mean is (low + lam*base + high)/(lam + 2).
    A degenerate row where low == high returns a point mass at that value.
    """
    if not low <= base <= high:
        raise ValueError(
            f"PERT requires low <= base <= high, got {low}, {base}, {high}"
        )
    if lam <= 0:
        raise ValueError(f"PERT shape must be positive, got {lam}")

    span = high - low
    if span == 0:
        # No spread to model. uniform with zero scale is a point mass at `low`.
        return stats.uniform(loc=low, scale=0)

    alpha = 1 + lam * (base - low) / span
    beta = 1 + lam * (high - base) / span
    return stats.beta(alpha, beta, loc=low, scale=span)


def pert_mean(low: float, base: float, high: float, lam: float = LAMBDA_DEFAULT) -> float:
    """The modified-PERT mean, stated directly as in the register documentation."""
    return (low + lam * base + high) / (lam + 2)


# ---------------------------------------------------------------------------------
# records
# ---------------------------------------------------------------------------------


@dataclass(frozen=True)
class Prior:
    """One register row, converted to engine units and frozen as a distribution."""

    param: str  # the register's "Model parameter"
    field: str  # where it lives in params.py, as "Model.field"
    row_id: int
    group: str
    assumption: str
    unit: str
    low: float  # engine units, after any percent conversion
    base: float
    high: float
    lam: float
    source_type: str
    confidence: str | None
    sens_rank: int | None
    scaled_by: float  # 1.0, or 100.0 for a percent row stored as a fraction
    dist: object  # scipy frozen distribution

    @property
    def mean(self) -> float:
        return pert_mean(self.low, self.base, self.high, self.lam)


@dataclass(frozen=True)
class SkippedRow:
    """A register row that yielded no distribution, and why."""

    row_id: int | None
    param: str
    source_type: str | None
    reason: str


@dataclass(frozen=True)
class Priors:
    """Priors for one register, plus what the register does not cover.

    ``by_param`` maps the register's "Model parameter" to a list of priors. It is always
    a list because the register genuinely repeats parameters -- six ``segments.discount``
    rows, one per payer segment, and two ``competitors.entry`` rows. Collapsing those to
    a single key would silently discard all but one.
    """

    by_param: dict[str, list[Prior]]
    skipped: list[SkippedRow]
    missing: dict[str, str]  # model field -> why it is required
    scope_mismatches: dict[str, str]
    register_path: Path

    def __len__(self) -> int:
        return sum(len(v) for v in self.by_param.values())

    def one(self, param: str) -> Prior:
        """The single prior for ``param``, or an error if the register repeats it."""
        found = self.by_param.get(param)
        if not found:
            raise KeyError(f"no register row for {param!r}")
        if len(found) > 1:
            raise KeyError(
                f"{param!r} has {len(found)} register rows "
                f"(ids {[p.row_id for p in found]}); use by_param[{param!r}]"
            )
        return found[0]

    def distributions(self) -> dict[str, list[object]]:
        """Just the frozen distributions, keyed as ``by_param``."""
        return {k: [p.dist for p in v] for k, v in self.by_param.items()}


class MissingRegisterRows(RuntimeError):
    """Raised when a model parameter has no register row behind it."""

    def __init__(self, missing: dict[str, str], priors: Priors):
        self.missing = missing
        self.priors = priors
        super().__init__(_format_missing(missing, priors))


def _format_missing(missing: dict[str, str], priors: Priors) -> str:
    lines = [
        f"{len(missing)} model parameter(s) have no usable row in "
        f"{priors.register_path.name}.",
        "",
        "The register is the source of truth for parameter values, so an unregistered "
        "parameter cannot be sampled -- it would sit at its default and report zero "
        "uncertainty, which is indistinguishable from a driver that does not matter.",
        "",
    ]
    for f, why in missing.items():
        lines.append(f"  {f:<34} {why}")
    lines += [
        "",
        "Add a register row per parameter (CLAUDE.md: new parameter -> register row "
        "first), or pass require_coverage=False to build the priors that do exist.",
    ]
    return "\n".join(lines)


# ---------------------------------------------------------------------------------
# reading
# ---------------------------------------------------------------------------------


def _camel_to_snake(name: str) -> str:
    """populationM -> population_m, bassP -> bass_p, persist12 -> persist12."""
    return re.sub(r"(?<=[a-z])(?=[A-Z])", "_", name).lower()


def _resolve_field(param: str) -> str | None:
    """Map a register "Model parameter" to "Model.field" in params.py, if it exists."""
    if param in EXPLICIT_FIELD_MAP:
        return EXPLICIT_FIELD_MAP[param]
    snake = _camel_to_snake(param)
    for model_name, model in _MODELS.items():
        if snake in model.model_fields:
            return f"{model_name}.{snake}"
    return None


def _is_percent(unit: str | None) -> bool:
    return "%" in (unit or "")


def _cell(ws, row: int, col: int):
    v = ws.cell(row, col).value
    return v.strip() if isinstance(v, str) else v


def _read_rows(path: Path) -> Iterable[dict]:
    ws = openpyxl.load_workbook(path, data_only=True)[SHEET]
    cols = {}
    for c in range(1, ws.max_column + 1):
        header = _cell(ws, HEADER_ROW, c)
        if header:
            cols[str(header)] = c

    required = ["ID", "Model parameter", "Unit", "Low", "Base", "High", "Source type"]
    absent = [c for c in required if c not in cols]
    if absent:
        raise ValueError(f"{path.name} is missing expected column(s): {absent}")

    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        row_id = _cell(ws, r, cols["ID"])
        if not isinstance(row_id, (int, float)):
            continue  # driver-group banner row, or blank
        yield {
            "row_id": int(row_id),
            "group": _cell(ws, r, cols.get("Driver group", 0)) if "Driver group" in cols else "",
            "assumption": _cell(ws, r, cols["Assumption"]) if "Assumption" in cols else "",
            "param": _cell(ws, r, cols["Model parameter"]),
            "unit": _cell(ws, r, cols["Unit"]),
            "low": _cell(ws, r, cols["Low"]),
            "base": _cell(ws, r, cols["Base"]),
            "high": _cell(ws, r, cols["High"]),
            "source_type": _cell(ws, r, cols["Source type"]),
            "confidence": _cell(ws, r, cols["Confidence"]) if "Confidence" in cols else None,
            "sens_rank": _cell(ws, r, cols["Sens. rank"]) if "Sens. rank" in cols else None,
        }


# ---------------------------------------------------------------------------------
# coverage
# ---------------------------------------------------------------------------------


def _model_fields() -> dict[str, str]:
    """Every parameter in params.py as "Model.field" -> a human label."""
    out = {}
    for model_name, model in _MODELS.items():
        for fname in model.model_fields:
            out[f"{model_name}.{fname}"] = fname
    return out


def _coverage(covered: set[str], skipped: list[SkippedRow]) -> dict[str, str]:
    """Model parameters with no usable register row behind them.

    A parameter whose row exists but was unusable reports that instead of "no register
    row" -- an unfilled row and an absent one need different fixes.
    """
    # Rows that named a real parameter but produced no distribution.
    unusable: dict[str, SkippedRow] = {}
    for s in skipped:
        resolved = _resolve_field(s.param) if s.param else None
        if resolved is not None and resolved not in covered:
            unusable.setdefault(resolved, s)

    missing = {}
    for field_path in _model_fields():
        if field_path in EXEMPT_FIELDS or field_path in covered:
            continue
        if field_path in unusable:
            s = unusable[field_path]
            missing[field_path] = f"register row {s.row_id} is unusable: {s.reason}"
        else:
            missing[field_path] = "no register row"
    return missing


# ---------------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------------


def load_priors(
    path: str | Path | None = None,
    *,
    require_coverage: bool = True,
) -> Priors:
    """Build modified-PERT priors from the assumptions register.

    Rows whose Source type is "Derived" are skipped -- they are computed by the model
    from other assumptions, so they carry no independent uncertainty.

    Raises ``MissingRegisterRows`` when a model parameter has no usable row, unless
    ``require_coverage`` is False.
    """
    register = Path(path) if path is not None else DEFAULT_REGISTER

    by_param: dict[str, list[Prior]] = {}
    skipped: list[SkippedRow] = []
    covered: set[str] = set()

    for row in _read_rows(register):
        param, src = row["param"], row["source_type"]

        if not param or param.startswith("("):
            skipped.append(
                SkippedRow(row["row_id"], str(param), src, "no model parameter named")
            )
            continue

        if src in SKIP_SOURCE_TYPES:
            skipped.append(
                SkippedRow(row["row_id"], param, src, f"source type is {src}")
            )
            continue

        values = [row["low"], row["base"], row["high"]]
        if any(not isinstance(v, (int, float)) for v in values):
            blank = [
                n for n, v in zip(("Low", "Base", "High"), values)
                if not isinstance(v, (int, float))
            ]
            skipped.append(
                SkippedRow(row["row_id"], param, src, f"blank {'/'.join(blank)}")
            )
            continue

        field_path = _resolve_field(param)
        if field_path is None:
            skipped.append(
                SkippedRow(
                    row["row_id"], param, src, "no matching parameter in params.py"
                )
            )
            continue

        # Percent rows are stored as fractions in the register; the model stores 28,
        # not 0.28. Convert here, and record the factor.
        scale = 100.0 if _is_percent(row["unit"]) else 1.0
        low, base, high = (float(v) * scale for v in values)

        lam = LAMBDA_ELICITED if src == "Elicited" else LAMBDA_DEFAULT

        try:
            dist = pert(low, base, high, lam)
        except ValueError as exc:
            skipped.append(SkippedRow(row["row_id"], param, src, str(exc)))
            continue

        rank = row["sens_rank"]
        by_param.setdefault(param, []).append(
            Prior(
                param=param,
                field=field_path,
                row_id=row["row_id"],
                group=str(row["group"] or ""),
                assumption=str(row["assumption"] or ""),
                unit=str(row["unit"] or ""),
                low=low,
                base=base,
                high=high,
                lam=lam,
                source_type=str(src),
                confidence=row["confidence"],
                sens_rank=int(rank) if isinstance(rank, (int, float)) else None,
                scaled_by=scale,
                dist=dist,
            )
        )
        covered.add(field_path)

    missing = _coverage(covered, skipped)
    priors = Priors(
        by_param=by_param,
        skipped=skipped,
        missing=missing,
        scope_mismatches={
            k: v for k, v in SCOPE_MISMATCHES.items() if k in by_param
        },
        register_path=register,
    )

    if missing and require_coverage:
        raise MissingRegisterRows(missing, priors)
    return priors
